"""Roster imports preserve existing teachers and can be rerun safely."""
from contextlib import closing
import json
import sqlite3

import pytest

from scripts.import_class_roster import apply_roster, load_roster


def roster(tmp_path):
    data = {"teacher": {"username": "belharet", "display_name": "Belharet Ali", "school_name": "Gana Ali", "subject": "Mathématiques"},
            "school_year": "2026/2027", "students": [{"class": "1M1", "name": "Eleve Un"}, {"class": "2M2", "name": "Eleve Deux"}]}
    path = tmp_path / "roster.json"
    path.write_text(json.dumps(data), encoding="utf-8")
    return path


def test_import_isolated_idempotent_and_backed_up(app, tmp_path):
    from core.db import DATABASE, get_db
    with app.app_context():
        db = get_db()
        db.execute("INSERT INTO users (username,password,nom_affichage,role,school_name) VALUES ('noura','existing-hash','Noura','prof','Autre ecole')")
        uid = db.execute("SELECT id FROM users WHERE username='noura'").fetchone()[0]
        db.execute("INSERT OR IGNORE INTO school_years (label,is_active) VALUES ('2025/2026',0)")
        db.execute("UPDATE school_years SET is_active=CASE WHEN label='2025/2026' THEN 1 ELSE 0 END")
        db.execute("INSERT INTO eleves (user_id,school_year,nom_complet,niveau) VALUES (?,'2025/2026','Eleve Noura','2M2')", (uid,))
        db.commit()
        before = tuple(db.execute("SELECT * FROM users WHERE id=?",(uid,)).fetchone())
    source = roster(tmp_path)
    access = tmp_path / "access.json"
    result = apply_roster(DATABASE, source, access)
    assert result["added"] == 2
    with closing(sqlite3.connect(result["backup"])) as backup:
        assert backup.execute("SELECT count(*) FROM users WHERE username='belharet'").fetchone()[0] == 0
    with closing(sqlite3.connect(DATABASE)) as db:
        assert tuple(db.execute("SELECT * FROM users WHERE id=?",(uid,)).fetchone()) == before
        assert db.execute("SELECT count(*) FROM eleves WHERE user_id=?",(uid,)).fetchone()[0] == 1
        assert db.execute("SELECT label FROM school_years WHERE is_active=1").fetchone()[0] == "2025/2026"
        assert db.execute("SELECT role,is_admin,lock_subject FROM users WHERE username='belharet'").fetchone() == ('prof',0,1)
        assert db.execute("SELECT count(*) FROM notes").fetchone()[0] == 0
        assert db.execute("SELECT count(*) FROM teacher_assignments WHERE user_id=?",(result["user_id"],)).fetchone()[0] == 2
        password_hash = db.execute("SELECT password FROM users WHERE username='belharet'").fetchone()[0]
    from werkzeug.security import check_password_hash
    credentials = json.loads(access.read_text(encoding="utf-8"))
    assert check_password_hash(password_hash,credentials["temporary_password"])
    assert apply_roster(DATABASE,source)["added"] == 0
    with closing(sqlite3.connect(DATABASE)) as db:
        assert db.execute("SELECT password FROM users WHERE username='belharet'").fetchone()[0] == password_hash
        assert db.execute("SELECT count(*) FROM eleves WHERE user_id=?",(result["user_id"],)).fetchone()[0] == 2


def test_duplicate_input_is_rejected(tmp_path):
    source = roster(tmp_path)
    data = json.loads(source.read_text())
    data["students"].append(data["students"][0])
    source.write_text(json.dumps(data))
    with pytest.raises(ValueError,match="duplicate"):
        load_roster(source)


def test_conflicting_username_is_not_modified(app,tmp_path):
    from core.db import DATABASE,get_db
    with app.app_context():
        db = get_db()
        db.execute("INSERT INTO users (username,password,nom_affichage,role,school_name) VALUES ('belharet','keep','Someone else','prof','Other')")
        db.commit()
    with pytest.raises(ValueError,match="different teacher"):
        apply_roster(DATABASE,roster(tmp_path))
    with closing(sqlite3.connect(DATABASE)) as db:
        assert db.execute("SELECT password FROM users WHERE username='belharet'").fetchone()[0] == 'keep'
        assert db.execute("SELECT count(*) FROM eleves").fetchone()[0] == 0


def test_credential_file_failure_rolls_back(app,tmp_path):
    from core.db import DATABASE
    access = tmp_path / "existing.json"
    access.write_text("keep")
    with pytest.raises(FileExistsError):
        apply_roster(DATABASE,roster(tmp_path),access)
    with closing(sqlite3.connect(DATABASE)) as db:
        assert db.execute("SELECT count(*) FROM users WHERE username='belharet'").fetchone()[0] == 0
    assert access.read_text() == "keep"

def test_excel_roster_reads_class_names_without_inventing_grades(tmp_path):
    from openpyxl import Workbook
    from scripts.import_class_roster import load_excel_roster
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "1M1"
    sheet.cell(8, 17, "Eleve Un")
    sheet.cell(8, 18, 1)
    sheet.cell(9, 17, "Eleve Deux")
    sheet.cell(9, 18, 2)
    source = tmp_path / "classes.xlsx"
    workbook.save(source)
    workbook.close()
    teacher = {"username": "math", "display_name": "Prof Math", "school_name": "Ecole", "subject": "Math"}
    data = load_excel_roster(source, teacher, "2026/2027")
    assert data["students"] == [{"class": "1M1", "name": "Eleve Un"}, {"class": "1M1", "name": "Eleve Deux"}]
    assert data["school_year"] == "2026/2027"
    assert len(data["source_sha256"]) == 64


def test_excel_roster_rejects_missing_student_number(tmp_path):
    from openpyxl import Workbook
    from scripts.import_class_roster import load_excel_roster
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "1M1"
    sheet.cell(8, 17, "Eleve Un")
    sheet.cell(8, 18, 2)
    source = tmp_path / "classes.xlsx"
    workbook.save(source)
    workbook.close()
    with pytest.raises(ValueError, match="numbering"):
        load_excel_roster(source, {}, "2026/2027")
